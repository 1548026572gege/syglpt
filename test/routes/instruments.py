# -*- coding: utf-8 -*-
"""仪器台账管理模块（3.2）：录入、查询筛选、状态更新、批量导入/导出"""
import io

from flask import Blueprint, request, jsonify, send_file
from openpyxl import Workbook, load_workbook

from models import db, Instrument
from auth import login_required, role_required, log_action

bp = Blueprint("instruments", __name__, url_prefix="/api/instruments")

VALID_STATUS = {"正常", "借出", "维修中", "损耗", "报废"}


@bp.get("")
@login_required
def list_instruments():
    kw = request.args.get("kw", "").strip()
    category = request.args.get("category", "").strip()
    status = request.args.get("status", "").strip()
    q = Instrument.query
    if kw:
        q = q.filter(db.or_(Instrument.name.contains(kw), Instrument.code.contains(kw),
                            Instrument.state_code.contains(kw)))
    if category:
        q = q.filter(Instrument.category == category)
    if status:
        q = q.filter(Instrument.status == status)
    items = q.order_by(Instrument.id.desc()).all()
    categories = [r[0] for r in db.session.query(Instrument.category)
                  .filter(Instrument.category.isnot(None)).distinct().all() if r[0]]
    # 当前筛选结果的总金额（单价 × 数量 之和，NULL 按 0 计）
    total_amount = q.with_entities(
        db.func.sum(db.func.coalesce(Instrument.price, 0)
                    * db.func.coalesce(Instrument.quantity, 0))).scalar() or 0
    return jsonify({"code": 0, "data": {"list": [i.to_dict() for i in items],
                                        "categories": categories,
                                        "total_amount": round(float(total_amount), 2)}})


@bp.post("")
@role_required("admin", "sysadmin")
def create_instrument():
    data = request.get_json(force=True, silent=True) or {}
    if not data.get("code") or not data.get("name"):
        return jsonify({"code": 400, "msg": "仪器编号和名称不能为空"}), 400
    if Instrument.query.filter_by(code=data["code"]).first():
        return jsonify({"code": 400, "msg": "仪器编号已存在"}), 400
    ins = Instrument(
        state_code=data.get("state_code", ""), code=data["code"], name=data["name"],
        model=data.get("model", ""), spec=data.get("spec", ""), unit=data.get("unit", ""),
        category=data.get("category", ""),
        quantity=int(data.get("quantity") or 1),
        price=float(data.get("price") or 0),
        requirement=data.get("requirement") or None,
        amount=float(data.get("amount") or 0),
        location=data.get("location", ""),
        status=data.get("status") if data.get("status") in VALID_STATUS else "正常",
        purchase_date=data.get("purchase_date", ""), remark=data.get("remark", ""))
    db.session.add(ins)
    db.session.commit()
    log_action("新增仪器", f"{ins.code} {ins.name}")
    return jsonify({"code": 0, "msg": "新增成功", "data": ins.to_dict()})


@bp.put("/<int:iid>")
@role_required("admin", "sysadmin")
def update_instrument(iid):
    ins = db.session.get(Instrument, iid)
    if not ins:
        return jsonify({"code": 404, "msg": "仪器不存在"}), 404
    data = request.get_json(force=True, silent=True) or {}
    for f in ("state_code", "name", "model", "spec", "unit", "category",
              "location", "purchase_date", "remark"):
        if f in data:
            setattr(ins, f, data[f])
    if "quantity" in data:
        ins.quantity = int(data["quantity"] or 1)
    if "price" in data:
        ins.price = float(data["price"] or 0)
    if "requirement" in data:
        ins.requirement = data.get("requirement") or None
    if "amount" in data:
        ins.amount = float(data["amount"] or 0)
    if data.get("status") in VALID_STATUS:
        ins.status = data["status"]
    db.session.commit()
    log_action("编辑仪器", f"{ins.code} {ins.name} 状态:{ins.status}")
    return jsonify({"code": 0, "msg": "保存成功", "data": ins.to_dict()})


@bp.delete("/<int:iid>")
@role_required("admin", "sysadmin")
def delete_instrument(iid):
    ins = db.session.get(Instrument, iid)
    if not ins:
        return jsonify({"code": 404, "msg": "仪器不存在"}), 404
    code, name = ins.code, ins.name
    db.session.delete(ins)
    db.session.commit()
    log_action("删除仪器", f"{code} {name}")
    return jsonify({"code": 0, "msg": "删除成功"})


@bp.post("/batch-delete")
@role_required("admin", "sysadmin")
def batch_delete_instruments():
    """批量删除勾选的仪器"""
    ids = (request.get_json(force=True, silent=True) or {}).get("ids") or []
    ids = [int(i) for i in ids if str(i).isdigit()]
    if not ids:
        return jsonify({"code": 400, "msg": "请先勾选要删除的仪器"}), 400
    items = Instrument.query.filter(Instrument.id.in_(ids)).all()
    n = len(items)
    info = "；".join(f"{i.code} {i.name}" for i in items[:10]) + ("…" if n > 10 else "")
    for i in items:
        db.session.delete(i)
    db.session.commit()
    log_action("批量删除仪器", f"共{n}台：{info}")
    return jsonify({"code": 0, "msg": f"已删除 {n} 台仪器"})


@bp.delete("/all")
@role_required("admin", "sysadmin")
def delete_all_instruments():
    """清空仪器台账（危险操作，记录日志）"""
    n = Instrument.query.count()
    if not n:
        return jsonify({"code": 400, "msg": "台账为空，无需删除"}), 400
    Instrument.query.delete()
    db.session.commit()
    log_action("清空仪器台账", f"共删除 {n} 台仪器")
    return jsonify({"code": 0, "msg": f"台账已清空，共删除 {n} 台仪器"})


HEADERS = ["国资编号", "分类代码", "器材名称", "规格", "单位", "数量", "单价",
           "存放位置", "状态", "金额", "备注"]


@bp.get("/export")
@login_required
def export_instruments():
    wb = Workbook()
    ws = wb.active
    ws.title = "仪器台账"
    ws.append(HEADERS)
    for i in Instrument.query.order_by(Instrument.id).all():
        ws.append([i.state_code, i.code, i.name, i.spec or i.model, i.unit, i.quantity, i.price,
                   i.location, i.status, i.amount, i.remark])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    log_action("导出仪器台账", f"共{Instrument.query.count()}条")
    return send_file(buf, as_attachment=True, download_name="仪器台账.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.get("/template")
@login_required
def import_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "仪器导入模板"
    ws.append(HEADERS)
    ws.append(["GZ20240001", "SW-YJ-001", "接种环", "标准型", "个", 30, 15.5,
               "生物实验室A-103", "正常", 465, ""])
    ws.append(["GZ20240002", "SW-YJ-002", "显微镜", "XSP-02，1600倍", "台", 50, 800,
               "生物实验室A-201", "正常", 40000, ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="仪器导入模板.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.post("/import")
@role_required("admin", "sysadmin")
def import_instruments():
    f = request.files.get("file")
    if not f or not f.filename.endswith((".xlsx", ".xlsm")):
        return jsonify({"code": 400, "msg": "请上传 xlsx 格式文件"}), 400
    try:
        ws = load_workbook(io.BytesIO(f.read())).active
    except Exception:
        return jsonify({"code": 400, "msg": "文件解析失败，请使用导入模板"}), 400
    ok, skip, errors = 0, 0, []
    used_codes = set()
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row):
            continue
        vals = [("" if v is None else str(v).strip()) for v in row] + [""] * (11 - len(row))
        # 表头：国资编号、分类代码、器材名称、规格、单位、数量、单价、存放位置、状态、金额、备注
        state_code, code, name, spec, unit, qty, price, location, status, amount, remark = vals[:11]
        # 只以器材名称作为导入依据，名称为空的行才跳过
        if not name:
            skip += 1
            errors.append(f"第{idx}行：器材名称为空，已跳过")
            continue
        # 分类代码为空或与已有/本文件内重复时，自动生成唯一编号，保证该行正常导入
        base_code = code or f"AUTO-{idx:04d}"
        new_code, seq = base_code, 2
        while new_code in used_codes or Instrument.query.filter_by(code=new_code).first():
            new_code = f"{base_code}-{seq}"
            seq += 1
        used_codes.add(new_code)
        try:
            qty = int(float(qty)) if qty else 1
        except ValueError:
            qty = 1
        try:
            price = float(price) if price else 0
        except ValueError:
            price = 0
        try:
            amount = float(amount) if amount else (price * qty)
        except ValueError:
            amount = price * qty
        db.session.add(Instrument(state_code=state_code, code=new_code, name=name, spec=spec, unit=unit,
                                  quantity=qty, price=price,
                                  amount=amount, location=location,
                                  status=status if status in VALID_STATUS else "正常",
                                  remark=remark))
        ok += 1
    db.session.commit()
    log_action("批量导入仪器", f"成功{ok}条 跳过{skip}条")
    return jsonify({"code": 0, "msg": f"导入完成：成功 {ok} 条，跳过 {skip} 条",
                    "data": {"ok": ok, "skip": skip, "errors": errors[:20]}})
