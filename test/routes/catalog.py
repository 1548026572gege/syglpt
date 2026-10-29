# -*- coding: utf-8 -*-
"""初中教学仪器配备目录模块：查看、按学科筛选、从配备目录 Excel 导入（预留配备标准）"""
import io

from flask import Blueprint, request, jsonify, send_file
from openpyxl import Workbook, load_workbook

from models import db, InstrumentCatalog
from auth import login_required, role_required, log_action

bp = Blueprint("catalog", __name__, url_prefix="/api/catalog")

# 表头列名 → 字段（自动去空格匹配；兼容"必配/选配"双列打√ 与 "必配或者选配"单列文本两种格式）
COL_ALIASES = {
    "name": ("器材名称",),
    "category": ("器材类型/学习主题",),
    "code": ("分类代码", "编码"),
    "old_code": ("老编码", "老编号", "原编码"),
    "spec": ("规格品名教学性能要求", "规格", "规格型号"),
    "unit": ("单位",),
    "quantity": ("数量",),
    "price": ("单价",),
    "require_combined": ("必配或者选配", "配备要求", "必配/选配", "是否必配"),
    "require_yes": ("必配",),
    "require_opt": ("选配",),
    "remark": ("备注",),
}


def _norm(v):
    return str(v).strip().replace(" ", "").replace("\u3000", "") if v is not None else ""


def _match_header(rows):
    """在前 6 行中定位表头行（含"器材名称"），返回 (header_row_idx, {field: col_idx})；
    "必配/选配"等子表头位于主表头下一行（合并单元格），一并识别"""
    for ri, row in enumerate(rows[:6]):
        cols = {}
        for ci, cell in enumerate(row):
            t = _norm(cell)
            if not t:
                continue
            for field, aliases in COL_ALIASES.items():
                if field not in cols and any(t == a or a in t for a in aliases):
                    cols[field] = ci
                    break
        if "name" in cols:
            if ri + 1 < len(rows):
                for ci, cell in enumerate(rows[ri + 1]):
                    t = _norm(cell)
                    for field in ("require_yes", "require_opt"):
                        if field not in cols and t and t in COL_ALIASES[field]:
                            cols[field] = ci
            return ri, cols
    return None, {}


def _cell(row, ci):
    if ci is None or ci >= len(row):
        return ""
    v = row[ci]
    return "" if v is None else (v if isinstance(v, (int, float)) else str(v).strip())


@bp.get("")
@login_required
def list_catalog():
    """配备目录列表：学科筛选 + 关键字搜索"""
    subject = request.args.get("subject", "").strip()
    kw = request.args.get("kw", "").strip()
    q = InstrumentCatalog.query
    if subject:
        q = q.filter(InstrumentCatalog.subject == subject)
    if kw:
        q = q.filter(db.or_(InstrumentCatalog.name.contains(kw),
                            InstrumentCatalog.code.contains(kw),
                            InstrumentCatalog.category.contains(kw)))
    items = q.order_by(InstrumentCatalog.subject, InstrumentCatalog.id).all()
    subjects = [r[0] for r in db.session.query(InstrumentCatalog.subject)
                .filter(InstrumentCatalog.subject.isnot(None)).distinct().all() if r[0]]
    return jsonify({"code": 0, "data": {"list": [i.to_dict() for i in items],
                                        "subjects": sorted(subjects)}})


@bp.get("/import-template")
@login_required
def import_template():
    """下载仪器配备目录导入模板（标准表头：分类代码/器材名称/规格/单位/数量/单价/必配或者选配/数量/金额/备注）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "仪器配备目录"
    headers = ["分类代码", "器材名称", "规格", "单位", "数量", "单价",
               "必配或者选配", "数量", "金额", "备注"]
    ws.append(headers)
    ws.append(["20201000101", "液晶投影机", "3LCD显示技术，亮度≥3300流明", "台", 1, 12800, "必配", "", "", ""])
    ws.append(["20201001101", "银幕", "白塑幕，亮度系数≥0.85", "幅", 1, 2500, "选配", "", "", ""])
    for i, w in enumerate([16, 24, 40, 8, 8, 10, 14, 8, 10, 20], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="仪器配备目录导入模板.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.post("/import")
@role_required("admin", "sysadmin")
def import_catalog():
    """导入配备目录 Excel（支持多学科 sheet，sheet 名作为学科名；同名学科数据先清空再导入）"""
    f = request.files.get("file")
    if not f or not f.filename.endswith((".xlsx", ".xlsm")):
        return jsonify({"code": 400, "msg": "请上传 xlsx 格式文件"}), 400
    try:
        wb = load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception:
        return jsonify({"code": 400, "msg": "文件解析失败"}), 400
    total, details = 0, []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hri, cols = _match_header(rows)
        if hri is None:
            details.append(f"{ws.title}：未找到含'器材名称'的表头行，已跳过")
            continue
        subject = ws.title.strip()
        InstrumentCatalog.query.filter_by(subject=subject).delete()
        n = 0
        last_cat = ""
        for row in rows[hri + 1:]:
            name = _cell(row, cols.get("name"))
            if not name or not str(name).strip():
                continue
            cat = _cell(row, cols.get("category"))
            if cat:
                last_cat = cat
            # 配备要求：优先取"必配或者选配"单列文本值，否则回退到必配/选配双列打√
            req = ""
            if cols.get("require_combined") is not None:
                v = _cell(row, cols["require_combined"])
                v = str(v).strip()
                if "必配" in v:
                    req = "必配"
                elif "选配" in v:
                    req = "选配"
            elif cols.get("require_yes") is not None and _cell(row, cols["require_yes"]):
                req = "必配"
            elif cols.get("require_opt") is not None and _cell(row, cols["require_opt"]):
                req = "选配"
            qty = _cell(row, cols.get("quantity"))
            price = _cell(row, cols.get("price"))
            db.session.add(InstrumentCatalog(
                subject=subject,
                category=last_cat or None,
                old_code=str(_cell(row, cols.get("old_code"))) or None,
                code=str(_cell(row, cols.get("code"))) or None,
                name=str(name).strip(),
                spec=_cell(row, cols.get("spec")) or None,
                unit=_cell(row, cols.get("unit")) or None,
                quantity=int(qty) if str(qty).strip().isdigit() else None,
                price=float(price) if str(price).replace(".", "").isdigit() else None,
                requirement=req or None,
                remark=_cell(row, cols.get("remark")) or None))
            n += 1
        total += n
        details.append(f"{subject}：导入 {n} 条")
    db.session.commit()
    log_action("导入仪器配备目录", f"共{total}条：" + "；".join(details))
    return jsonify({"code": 0, "msg": f"导入完成，共 {total} 条", "data": {"details": details}})


@bp.delete("/<int:cid>")
@role_required("admin", "sysadmin")
def delete_catalog_item(cid):
    item = db.session.get(InstrumentCatalog, cid)
    if not item:
        return jsonify({"code": 404, "msg": "记录不存在"}), 404
    db.session.delete(item)
    db.session.commit()
    log_action("删除配备目录记录", f"{item.subject} {item.name}")
    return jsonify({"code": 0, "msg": "删除成功"})


@bp.post("/batch-delete")
@role_required("admin", "sysadmin")
def batch_delete_catalog():
    """批量删除配备目录记录（按 id 列表，一次完成，解决逐条删除卡顿）"""
    ids = (request.get_json(force=True, silent=True) or {}).get("ids") or []
    ids = [int(i) for i in ids if str(i).isdigit()]
    if not ids:
        return jsonify({"code": 400, "msg": "请选择要删除的记录"}), 400
    n = InstrumentCatalog.query.filter(InstrumentCatalog.id.in_(ids)).count()
    InstrumentCatalog.query.filter(InstrumentCatalog.id.in_(ids)).delete(synchronize_session=False)
    db.session.commit()
    log_action("批量删除配备目录", f"共{n}条")
    return jsonify({"code": 0, "msg": f"已删除 {n} 条"})
