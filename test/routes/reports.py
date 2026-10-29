# -*- coding: utf-8 -*-
"""实验登记与报告模块（3.4）：使用登记、过程记录、报告单生成、导出、附件上传"""
import io
import os
from datetime import datetime

from flask import Blueprint, request, jsonify, current_app, send_file, g
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename

from models import db, ExperimentReport, Reservation
from auth import login_required, role_required, log_action

bp = Blueprint("reports", __name__, url_prefix="/api/reports")


def _parse_dt(s):
    if not s:
        return None
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


@bp.get("")
@login_required
def list_reports():
    kw = request.args.get("kw", "").strip()
    q = ExperimentReport.query
    if g.current_user.role == "teacher":
        q = q.filter(ExperimentReport.teacher_id == g.current_user.id)
    if kw:
        q = q.filter(ExperimentReport.experiment_name.contains(kw))
    items = q.order_by(ExperimentReport.id.desc()).all()
    return jsonify({"code": 0, "data": [r.to_dict() for r in items]})


@bp.get("/<int:rid>")
@login_required
def get_report(rid):
    r = db.session.get(ExperimentReport, rid)
    if not r:
        return jsonify({"code": 404, "msg": "报告不存在"}), 404
    if g.current_user.role == "teacher" and r.teacher_id != g.current_user.id:
        return jsonify({"code": 403, "msg": "无权限查看"}), 403
    return jsonify({"code": 0, "data": r.to_dict()})


@bp.post("")
@role_required("teacher", "admin", "sysadmin")
def create_report():
    data = request.get_json(force=True, silent=True) or {}
    if not data.get("experiment_name"):
        return jsonify({"code": 400, "msg": "实验名称不能为空"}), 400
    reservation_id = data.get("reservation_id") or None
    if reservation_id:
        rv = db.session.get(Reservation, int(reservation_id))
        if not rv:
            return jsonify({"code": 404, "msg": "关联预约不存在"}), 404
    r = ExperimentReport(
        reservation_id=reservation_id, teacher_id=g.current_user.id,
        experiment_name=data["experiment_name"],
        actual_time=_parse_dt(data.get("actual_time")) or datetime.now(),
        class_name=data.get("class_name", ""), group_info=data.get("group_info", ""),
        instruments_used=data.get("instruments_used", ""),
        content=data.get("content", ""), result=data.get("result", ""),
        exception=data.get("exception", ""))
    db.session.add(r)
    db.session.commit()
    log_action("登记实验报告", r.experiment_name)
    return jsonify({"code": 0, "msg": "报告已保存", "data": r.to_dict()})


@bp.put("/<int:rid>")
@login_required
def update_report(rid):
    r = db.session.get(ExperimentReport, rid)
    if not r:
        return jsonify({"code": 404, "msg": "报告不存在"}), 404
    if g.current_user.role == "teacher" and r.teacher_id != g.current_user.id:
        return jsonify({"code": 403, "msg": "只能编辑自己的报告"}), 403
    data = request.get_json(force=True, silent=True) or {}
    for f in ("experiment_name", "class_name", "group_info", "instruments_used",
              "content", "result", "exception"):
        if f in data:
            setattr(r, f, data[f])
    if "actual_time" in data:
        r.actual_time = _parse_dt(data["actual_time"]) or r.actual_time
    db.session.commit()
    return jsonify({"code": 0, "msg": "保存成功", "data": r.to_dict()})


@bp.delete("/<int:rid>")
@login_required
def delete_report(rid):
    r = db.session.get(ExperimentReport, rid)
    if not r:
        return jsonify({"code": 404, "msg": "报告不存在"}), 404
    if g.current_user.role == "teacher" and r.teacher_id != g.current_user.id:
        return jsonify({"code": 403, "msg": "只能删除自己的报告"}), 403
    db.session.delete(r)
    db.session.commit()
    log_action("删除实验报告", f"#{rid} {r.experiment_name}")
    return jsonify({"code": 0, "msg": "删除成功"})


@bp.post("/<int:rid>/attachment")
@login_required
def upload_attachment(rid):
    r = db.session.get(ExperimentReport, rid)
    if not r:
        return jsonify({"code": 404, "msg": "报告不存在"}), 404
    if g.current_user.role == "teacher" and r.teacher_id != g.current_user.id:
        return jsonify({"code": 403, "msg": "无权限"}), 403
    f = request.files.get("file")
    if not f:
        return jsonify({"code": 400, "msg": "请选择文件"}), 400
    ext = os.path.splitext(f.filename)[1].lower()
    if ext not in (".pdf", ".doc", ".docx", ".xls", ".xlsx", ".png", ".jpg", ".jpeg", ".zip"):
        return jsonify({"code": 400, "msg": "不支持的文件类型"}), 400
    fname = f"report_{rid}_{int(datetime.now().timestamp())}{ext}"
    f.save(os.path.join(current_app.config["UPLOAD_FOLDER"], secure_filename(fname)))
    r.attachment = fname
    db.session.commit()
    return jsonify({"code": 0, "msg": "附件上传成功", "data": {"attachment": fname}})


@bp.get("/<int:rid>/export")
@login_required
def export_report(rid):
    """导出标准化实验报告单（xlsx）"""
    r = db.session.get(ExperimentReport, rid)
    if not r:
        return jsonify({"code": 404, "msg": "报告不存在"}), 404
    if g.current_user.role == "teacher" and r.teacher_id != g.current_user.id:
        return jsonify({"code": 403, "msg": "无权限导出"}), 403
    wb = Workbook()
    ws = wb.active
    ws.title = "实验报告单"
    rows = [
        ["生物实验报告单"],
        ["实验名称", r.experiment_name],
        ["任课教师", r.teacher.name if r.teacher else ""],
        ["实验时间", r.actual_time.strftime("%Y-%m-%d %H:%M") if r.actual_time else ""],
        ["参与班级", r.class_name or ""],
        ["分组情况", r.group_info or ""],
        ["使用仪器", r.instruments_used or ""],
        ["实验内容与过程", r.content or ""],
        ["实验结果", r.result or ""],
        ["异常情况", r.exception or ""],
    ]
    for row in rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 70
    from openpyxl.styles import Font, Alignment
    ws["A1"].font = Font(bold=True, size=14)
    for i in range(2, len(rows) + 1):
        ws.cell(row=i, column=1).font = Font(bold=True)
        ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    log_action("导出实验报告单", f"#{rid} {r.experiment_name}")
    return send_file(buf, as_attachment=True,
                     download_name=f"实验报告单_{r.experiment_name}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.get("/my-reservations")
@login_required
def my_finished_reservations():
    """教师可用于关联报告的已通过预约"""
    q = Reservation.query.filter(Reservation.teacher_id == g.current_user.id,
                                 Reservation.status.in_(["已通过", "已完成"])) \
        .order_by(Reservation.start_time.desc())
    return jsonify({"code": 0, "data": [r.to_dict() for r in q.all()]})


# ---------------- Excel 批量导入 / 导出 ----------------

IMPORT_HEADERS = ["实验名称*", "实际使用时间", "参与班级", "分组情况",
                  "使用仪器", "实验内容与过程", "实验结果", "异常情况"]
EXPORT_HEADERS = ["序号", "实验名称", "任课教师", "实际使用时间", "参与班级",
                  "分组情况", "使用仪器", "实验内容与过程", "实验结果", "异常情况", "登记时间"]

# 自行导入：按表头列名自动识别（列顺序不限，支持常见别名，空格自动去除）
FIELD_ALIASES = {
    "experiment_name": ("实验名称", "实验名称*", "实验名", "名称", "实验项目"),
    "actual_time": ("实际使用时间", "实验时间", "使用时间", "实际时间", "日期", "时间", "日期时间"),
    "class_name": ("参与班级", "班级", "上课班级", "授课班级", "任教班级"),
    "group_info": ("分组情况", "分组", "分组方式", "分组信息"),
    "instruments_used": ("使用仪器", "实验仪器", "仪器", "仪器用品", "所需仪器", "实验仪器、用品", "实验仪器、用品"),
    "content": ("实验内容与过程", "实验内容", "内容", "实验过程", "过程"),
    "result": ("实验结果", "结果", "实验结论", "结论"),
    "exception": ("异常情况", "异常", "备注", "问题", "异常与处理"),
}


def _match_columns(header):
    """按表头文字自动映射列位置 → 标准字段，返回 {field: col_index}"""
    mapping = {}
    for idx, cell in enumerate(header):
        if cell is None:
            continue
        name = str(cell).strip().replace(" ", "")
        if not name:
            continue
        for field, aliases in FIELD_ALIASES.items():
            if field not in mapping and name in [a.replace(" ", "") for a in aliases]:
                mapping[field] = idx
                break
    return mapping


@bp.get("/import-template")
@login_required
def import_template():
    """下载实验报告导入模板（任课教师填写后上传）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "实验报告导入"
    ws.append(IMPORT_HEADERS)
    ws.append(["观察洋葱表皮细胞结构", "2026-10-28 08:00", "高一(1)班", "10组，每组4人",
               "显微镜10台、镊子10把、碘液", "制作洋葱表皮临时装片并观察细胞结构",
               "学生能观察到清晰的细胞轮廓", "无"])
    ws.append(["观察植物细胞的质壁分离", "2026-10-29 10:00", "高一(2)班", "12组，每组4人",
               "显微镜12台、蔗糖溶液", "制作紫洋葱表皮装片，滴加蔗糖溶液观察质壁分离",
               "观察到质壁分离现象", "个别组装片气泡较多"])
    ws.column_dimensions["A"].width = 24
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="实验报告导入模板.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.post("/import")
@role_required("teacher", "admin", "sysadmin")
def import_reports():
    """任课教师导入实验报告单：支持标准模板或自行整理的 Excel（按表头列名自动识别，实验名称列必填）"""
    f = request.files.get("file")
    if not f or not f.filename.endswith((".xlsx", ".xlsm")):
        return jsonify({"code": 400, "msg": "请上传 xlsx 格式文件"}), 400
    try:
        ws = load_workbook(io.BytesIO(f.read())).active
    except Exception:
        return jsonify({"code": 400, "msg": "文件解析失败，请使用导入模板"}), 400
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return jsonify({"code": 400, "msg": "文件内容为空"}), 400
    colmap = _match_columns(rows[0])
    if "experiment_name" not in colmap:
        return jsonify({"code": 400,
                        "msg": "未识别到表头中的“实验名称”列，请下载模板填写后导入，或将列名改为“实验名称”"}), 400
    ok, skip, errors = 0, 0, []

    def cell(vals, field):
        idx = colmap.get(field)
        if idx is None or idx >= len(vals):
            return ""
        v = vals[idx]
        return "" if v is None else (v if isinstance(v, datetime) else str(v).strip())

    for idx, row in enumerate(rows[1:], start=2):
        if not row or not any(row):
            continue
        name = cell(row, "experiment_name")
        if not name:
            skip += 1
            errors.append(f"第{idx}行：缺少实验名称，已跳过")
            continue
        atime = cell(row, "actual_time")
        if isinstance(atime, datetime):
            actual_time = atime
        else:
            actual_time = _parse_dt(atime) or datetime.now()
        db.session.add(ExperimentReport(
            teacher_id=g.current_user.id, experiment_name=name, actual_time=actual_time,
            class_name=cell(row, "class_name"), group_info=cell(row, "group_info"),
            instruments_used=cell(row, "instruments_used"), content=cell(row, "content"),
            result=cell(row, "result"), exception=cell(row, "exception")))
        ok += 1
    db.session.commit()
    log_action("批量导入实验报告", f"成功{ok}条 跳过{skip}条")
    return jsonify({"code": 0, "msg": f"导入完成：成功 {ok} 条，跳过 {skip} 条",
                    "data": {"ok": ok, "skip": skip, "errors": errors[:20]}})


@bp.get("/export")
@login_required
def export_reports():
    """导出实验报告汇总表（教师导出自己的，管理员导出全部）"""
    q = ExperimentReport.query
    if g.current_user.role == "teacher":
        q = q.filter(ExperimentReport.teacher_id == g.current_user.id)
    items = q.order_by(ExperimentReport.id.desc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "实验报告汇总"
    ws.append(EXPORT_HEADERS)
    for idx, r in enumerate(items, start=1):
        ws.append([idx, r.experiment_name, r.teacher.name if r.teacher else "",
                   r.actual_time.strftime("%Y-%m-%d %H:%M") if r.actual_time else "",
                   r.class_name or "", r.group_info or "", r.instruments_used or "",
                   r.content or "", r.result or "", r.exception or "",
                   r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else ""])
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 12
    for col in "DEFGH":
        ws.column_dimensions[col].width = 20
    for col in "IJK":
        ws.column_dimensions[col].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    log_action("导出实验报告汇总", f"共{len(items)}条")
    return send_file(buf, as_attachment=True, download_name="实验报告汇总表.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
